"""Exportação da base final em CSV, Excel (multi-abas) e JSON."""
import csv
import json
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).parent.parent
FINAL_DIR = ROOT / "data" / "final"

# Paleta do projeto
COR_ACAI     = "3B1E54"   # roxo açaí — cabeçalho principal
COR_VERDE    = "2E7D32"   # verde amazônico — abas secundárias
COR_PARA_AZ  = "003DA5"   # azul Pará — destaque
COR_NEUTRO   = "E5E5E5"   # linhas alternadas
COR_BRANCO   = "FFFFFF"
COR_TEXTO_C  = "FFFFFF"   # texto em fundo escuro
COR_TEXTO_E  = "1E1E1E"   # texto em fundo claro


def _get_fieldnames(rows):
    return list(rows[0].keys()) if rows else []


def export_csv(rows: list, filename: str = "acai_estabelecimentos.csv"):
    FINAL_DIR.mkdir(parents=True, exist_ok=True)
    path = FINAL_DIR / filename
    fields = _get_fieldnames(rows)
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    print(f"  [exporter] CSV salvo: {path} ({len(rows)} registros)")
    return path


def _header_style(ws, row_num, fields, bg_color, font_color=COR_TEXTO_C):
    from openpyxl.styles import Font, PatternFill, Alignment
    fill = PatternFill("solid", fgColor=bg_color)
    font = Font(color=font_color, bold=True, size=10)
    for col_idx, field in enumerate(fields, 1):
        cell = ws.cell(row=row_num, column=col_idx, value=field)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)


def _write_rows(ws, rows, fields, start_row=2):
    from openpyxl.styles import PatternFill, Alignment
    fill_alt = PatternFill("solid", fgColor=COR_NEUTRO)
    for r_idx, row in enumerate(rows, start_row):
        for c_idx, field in enumerate(fields, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=row.get(field, ""))
            cell.alignment = Alignment(wrap_text=False)
            if r_idx % 2 == 0:
                cell.fill = fill_alt


def _auto_width(ws, fields, rows, max_w=60):
    from openpyxl.utils import get_column_letter
    for c_idx, field in enumerate(fields, 1):
        col_max = max(
            len(str(field)),
            max((len(str(row.get(field, "") or "")) for row in rows), default=0)
        )
        ws.column_dimensions[get_column_letter(c_idx)].width = min(col_max + 2, max_w)


def export_excel(
    rows_final: list,
    rows_controle: list,
    rows_duplicates: list,
    rows_bairros: list,
    resumo: dict,
    filename: str = "acai_estabelecimentos.xlsx",
):
    FINAL_DIR.mkdir(parents=True, exist_ok=True)
    path = FINAL_DIR / filename
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = openpyxl.Workbook()

        # ── ABA 1: Estabelecimentos ──────────────────────────────────────
        ws1 = wb.active
        ws1.title = "Estabelecimentos"
        fields1 = _get_fieldnames(rows_final)
        _header_style(ws1, 1, fields1, COR_ACAI)
        _write_rows(ws1, rows_final, fields1)
        _auto_width(ws1, fields1, rows_final)
        ws1.freeze_panes = "A2"
        ws1.auto_filter.ref = ws1.dimensions

        # ── ABA 2: Controle_Coleta ───────────────────────────────────────
        ws2 = wb.create_sheet("Controle_Coleta")
        fields2 = _get_fieldnames(rows_controle)
        if not fields2:
            fields2 = ["bairro", "termo", "consulta_completa", "status",
                       "inicio", "fim", "quantidade_resultados", "job_id", "erro"]
        _header_style(ws2, 1, fields2, COR_VERDE)
        _write_rows(ws2, rows_controle, fields2)
        _auto_width(ws2, fields2, rows_controle)
        ws2.freeze_panes = "A2"

        # ── ABA 3: Duplicados ────────────────────────────────────────────
        ws3 = wb.create_sheet("Duplicados")
        fields3 = _get_fieldnames(rows_duplicates)
        if not fields3:
            fields3 = ["place_id_principal", "regra_utilizada",
                       "registro_mantido_title", "registro_descartado_title", "motivo"]
        _header_style(ws3, 1, fields3, COR_PARA_AZ)
        _write_rows(ws3, rows_duplicates, fields3)
        _auto_width(ws3, fields3, rows_duplicates)
        ws3.freeze_panes = "A2"

        # ── ABA 4: Bairros ───────────────────────────────────────────────
        ws4 = wb.create_sheet("Bairros")
        fields4 = _get_fieldnames(rows_bairros)
        if not fields4:
            fields4 = ["bairro", "latitude", "longitude", "geocodificacao_status",
                       "fonte", "data_geocodificacao"]
        _header_style(ws4, 1, fields4, COR_ACAI)
        _write_rows(ws4, rows_bairros, fields4)
        _auto_width(ws4, fields4, rows_bairros)

        # ── ABA 5: Resumo ────────────────────────────────────────────────
        ws5 = wb.create_sheet("Resumo")
        ws5.column_dimensions["A"].width = 35
        ws5.column_dimensions["B"].width = 20

        # Título
        ws5["A1"] = "AÇAÍ MAP BELÉM — Resumo da Coleta"
        ws5["A1"].font = Font(color=COR_ACAI, bold=True, size=14)
        ws5.merge_cells("A1:B1")

        row_n = 3
        fill_h = PatternFill("solid", fgColor=COR_ACAI)
        font_h = Font(color=COR_BRANCO, bold=True)
        fill_v = PatternFill("solid", fgColor=COR_NEUTRO)

        for k, v in resumo.items():
            label = k.replace("_", " ").title()
            cell_k = ws5.cell(row=row_n, column=1, value=label)
            cell_v = ws5.cell(row=row_n, column=2, value=v)
            if row_n % 2 == 0:
                cell_k.fill = fill_v
                cell_v.fill = fill_v
            row_n += 1

        wb.save(path)
        print(f"  [exporter] Excel salvo: {path} (5 abas, {len(rows_final)} registros)")
    except ImportError:
        print("  [exporter] AVISO: openpyxl não instalado. Excel não gerado.")
    return path


def export_resumo(resumo: dict, filename: str = "resumo_coleta.json"):
    FINAL_DIR.mkdir(parents=True, exist_ok=True)
    path = FINAL_DIR / filename
    with open(path, "w", encoding="utf-8") as f:
        json.dump(resumo, f, ensure_ascii=False, indent=2)
    print(f"  [exporter] Resumo salvo: {path}")
    return path


def build_resumo(
    bairros_processados, bairros_concluidos, consultas_planejadas,
    consultas_executadas, consultas_concluidas, consultas_com_erro,
    resultados_brutos, registros_tratados, duplicidades,
    registros_finais, sem_coordenadas, sem_place_id, sem_avaliacao,
) -> dict:
    return {
        "projeto": "Açaí Map Belém",
        "data_coleta": datetime.now().strftime("%Y-%m-%d"),
        "hora_geracao": datetime.now().strftime("%H:%M:%S"),
        "bairros": bairros_processados,
        "bairros_concluidos": bairros_concluidos,
        "consultas_planejadas": consultas_planejadas,
        "consultas_executadas": consultas_executadas,
        "consultas_concluidas": consultas_concluidas,
        "consultas_com_erro": consultas_com_erro,
        "resultados_brutos": resultados_brutos,
        "registros_tratados": registros_tratados,
        "duplicidades": duplicidades,
        "registros_finais": registros_finais,
        "sem_coordenadas": sem_coordenadas,
        "sem_place_id": sem_place_id,
        "sem_avaliacao": sem_avaliacao,
    }


def load_csv_safe(path) -> list:
    """Carrega CSV com tratamento de arquivo vazio ou inexistente."""
    p = Path(path)
    if not p.exists() or p.stat().st_size == 0:
        return []
    with open(p, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def load_bairros_coordenadas() -> list:
    path = ROOT / "data" / "reference" / "bairros_coordenadas.csv"
    return load_csv_safe(path)
